"""
The equipment register as a workbook made to be read, not just filtered.

Three sheets:

* **Equipment** — one row per asset: its photo, IDs, description, state,
  where it is in use (station, customer, process, models) and its calibration
  and maintenance dates, under coloured group headings.
* **Stations** — station by station: the machines it shares, then each model
  run on it with the machines that model brings, a missing one shown as such.
* **Summary** — per equipment type: how many, how many in use, and how many
  calibrations are overdue.

Dates already past are red, those within 30 days amber; an asset that cannot be
used says why in red ("Can be used", as of the export day). The colours and
the whole Summary are Excel formulas against today's date, so they follow
the Equipment sheet as it is edited.
"""

from datetime import date, timedelta

from django.conf import settings
from django.db.models import Prefetch
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .services import PHOTO_CELL_PX, photo_jpeg
from .xlsx_images import put_in_cell

FONT = "Calibri"
DARK = "1F3864"
LINE = Side(style="thin", color="D0D7E2")
BORDER = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)
ZEBRA = PatternFill("solid", fgColor="F5F7FA")
RED, AMBER, GREEN, GREY = "C00000", "B26A00", "2E7D32", "7F8C8D"
ROW_PX = PHOTO_CELL_PX + 8

# (group, colour, [(heading, width)])
EQUIPMENT_GROUPS = [
    ("Asset", "DCE6F1", [("Photo", 11), ("Equipment No", 14), ("ZT Code", 13), ("Asset No.", 17), ("S/N", 18)]),
    ("Description", "E2EFDA", [("Name", 36), ("Type", 18), ("Manufacturer", 15), ("Model", 16)]),
    ("Where & state", "FFF2CC", [("Plant", 7), ("Status", 11), ("Label", 8), ("Can be used", 24)]),
    ("In use at", "FCE4D6", [("Station", 14), ("Customer", 11), ("Process", 9), ("Used as", 16), ("Models", 24)]),
    ("Calibration", "EDE2F6", [("Needs", 7), ("Every (months)", 9), ("Last", 11), ("Due", 11)]),
    ("Maintenance", "DDEBF7", [("Every (days)", 9), ("Last", 11), ("Due", 11)]),
    ("Other", "EDEDED", [("Specifications", 30), ("Notes", 46)]),
]

LAST_ROW = 2000          # formulas reach this far, so rows added by hand count too


# Pictures go *in* the photo cell (save with xlsx_images.save_workbook), so
# they move with their row when the sheet is sorted or filtered. Kept larger
# than the cell: Excel shrinks them to fit, and they stay sharp when the row
# is made taller.
PHOTO_STORED_PX = 240


def _photo(sheet, cell_ref, asset):
    jpeg = photo_jpeg(asset.main_photo, PHOTO_STORED_PX)
    if jpeg:
        put_in_cell(sheet, cell_ref, jpeg)


def _cell(sheet, row, column, value=None, *, bold=False, colour=None, size=10, fill=None, wrap=True,
          align="left", border=True, number_format=None):
    cell = sheet.cell(row=row, column=column, value=value)
    cell.font = Font(name=FONT, size=size, bold=bold, color=colour)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        cell.fill = fill if isinstance(fill, PatternFill) else PatternFill("solid", fgColor=fill)
    if border:
        cell.border = BORDER
    if number_format:
        cell.number_format = number_format
    return cell


def _due_colour(due, today):
    if not due:
        return None
    if due == "Unknown":          # needs calibration, no date on record
        return RED
    if due < today:
        return RED
    if due <= today + timedelta(days=30):
        return AMBER
    return None


def _first_due(plans):
    plans = [plan for plan in plans if plan.is_active]
    return min(plans, key=lambda plan: plan.next_due_at or date.max, default=None)


def _title(sheet, text, subtitle, width):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    _cell(sheet, 1, 1, text, bold=True, size=15, colour=DARK, border=False, wrap=False)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    _cell(sheet, 2, 1, subtitle, size=9, colour=GREY, border=False, wrap=False)
    sheet.row_dimensions[1].height = 24


def _usage(asset, station_models):
    """(station, customer, process, used as, models) for an asset's open fitting."""
    fitting = next((inst for inst in asset.open_fittings), None)
    if fitting is None:
        return None, "", "", "", ""
    if fitting.station_id:
        station = fitting.station
        return station, station.customer.name if station.customer else "", station.process.code, "Shared", \
            ", ".join(station_models.get(station.pk, []))
    test_set = fitting.test_set
    station = test_set.station if test_set else None
    used_as = fitting.slot.category.name if fitting.slot_id else "Own machine"
    product = test_set.product_process.product if test_set else None
    # A model's own machine may serve several models of its station: name them all.
    codes = [inst.test_set.product_process.product.code for inst in asset.open_fittings if inst.test_set_id]
    return (
        station,
        product.customer.name if product and product.customer else "",
        test_set.product_process.process.code if test_set else "",
        used_as,
        ", ".join(dict.fromkeys(codes)),
    )


def _equipment_sheet(sheet, assets, station_models, today, subtitle):
    columns = [heading for _, _, cols in EQUIPMENT_GROUPS for heading, _ in cols]
    width = len(columns)
    _title(sheet, "Equipment register", subtitle, width)

    # Group bands (row 3), then the column headings (row 4).
    column = 1
    for group, colour, cols in EQUIPMENT_GROUPS:
        if len(cols) > 1:
            sheet.merge_cells(start_row=3, start_column=column, end_row=3, end_column=column + len(cols) - 1)
        for offset in range(len(cols)):
            _cell(sheet, 3, column + offset, group if offset == 0 else None, bold=True, fill=colour,
                  colour=DARK, align="center")
        for heading, col_width in cols:
            _cell(sheet, 4, column, heading, bold=True, colour="FFFFFF", fill=DARK, align="center")
            sheet.column_dimensions[get_column_letter(column)].width = col_width
            column += 1
    sheet.row_dimensions[3].height = 18
    sheet.row_dimensions[4].height = 30

    # Where each column is, by group: "Last" and "Due" come twice.
    at = {}
    column = 1
    for group, _, cols in EQUIPMENT_GROUPS:
        for heading, _ in cols:
            at[(group, heading)] = get_column_letter(column)
            column += 1

    row = 5
    for index, asset in enumerate(assets):
        station, customer, process, used_as, models = _usage(asset, station_models)
        calibration = _first_due(asset.calibration_plans.all())
        maintenance = _first_due(asset.maintenance_plans.all())
        unknown = asset.calibration_date_unknown
        values = [
            None, asset.equipment_no, asset.code, asset.asset_no, asset.serial_number,
            asset.name, asset.category.name, asset.manufacturer, asset.model_number,
            asset.plant_display if asset.plant else "", asset.get_status_display(),
            "Yes" if asset.labelled else "No", _can_be_used(asset, calibration, maintenance, today),
            station.code if station else "", customer, process, used_as, models,
            "Yes" if asset.is_measuring_device else "No",
            calibration.interval_months if calibration else None,
            "Unknown" if unknown else (calibration.last_calibrated_at if calibration else None),
            "Unknown" if unknown else (calibration.next_due_at if calibration else None),
            maintenance.interval_days if maintenance else None,
            maintenance.last_done_at if maintenance else None,
            maintenance.next_due_at if maintenance else None,
            asset.specifications, asset.notes,
        ]
        zebra = ZEBRA if index % 2 else None
        for column, value in enumerate(values, start=1):
            heading = columns[column - 1]
            bold, number_format, align = False, None, "left"
            if heading in ("Last", "Due"):
                number_format, align = "dd/mm/yyyy", "center"
            elif heading in ("Every (months)", "Every (days)", "Needs", "Plant", "Process", "Label", "Status"):
                align = "center"
            elif heading in ("Equipment No", "Can be used"):
                bold = True
            elif heading == "Station":
                bold = bool(value)
            _cell(sheet, row, column, value, bold=bold, fill=zebra, number_format=number_format, align=align)
        _photo(sheet, f"A{row}", asset)
        sheet.row_dimensions[row].height = ROW_PX * 0.75
        row += 1

    _equipment_colours(sheet, at)
    sheet.freeze_panes = "C5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(width)}{max(row - 1, 4)}"
    _print_setup(sheet, 4)
    return at


def _can_be_used(asset, calibration, maintenance, today):
    """Whether the asset can be used on the export day, in plain words (no day counts)."""
    if not asset.is_available:
        return asset.get_status_display()
    if asset.calibration_date_unknown:
        return "Calibration date unknown"
    if calibration and calibration.next_due_at and calibration.next_due_at < today:
        return "Calibration overdue"
    if maintenance and maintenance.next_due_at and maintenance.next_due_at < today:
        return "Maintenance overdue"
    return "OK"


def _equipment_colours(sheet, at):
    """Colours that follow the data and the date, as conditional formats."""
    def rule(columns, formula, colour, fill=None):
        for letter in columns:
            sheet.conditional_formatting.add(
                f"{letter}5:{letter}{LAST_ROW}",
                FormulaRule(formula=[formula.replace("@", f"{letter}5")], stopIfTrue=True,
                            font=Font(color=colour, bold=True),
                            fill=PatternFill("solid", fgColor=fill) if fill else None),
            )

    dues = [at[("Calibration", "Due")], at[("Maintenance", "Due")]]
    rule(dues + [at[("Calibration", "Last")]], '@="Unknown"', RED, "FDE2E1")
    rule(dues, "AND(ISNUMBER(@),@<TODAY())", RED, "FDE2E1")
    rule(dues, "AND(ISNUMBER(@),@<=TODAY()+30)", AMBER, "FFF2CC")
    can = at[("Where & state", "Can be used")]
    rule([can], '@="OK"', GREEN)
    rule([can], 'AND(@<>"OK",@<>"")', RED, "FDE2E1")
    status = at[("Where & state", "Status")]
    rule([status], '@="Broken"', RED, "FDE2E1")
    rule([status], '@="Spare"', GREY)
    rule([status], '@="In use"', GREEN)


def _stations_sheet(sheet, stations, today, subtitle):
    from systems.services import evaluate_set, unusable_reason

    headings = [("Photo", 11), ("Used as", 24), ("Equipment No", 14), ("ZT Code", 13), ("Name", 36),
                ("Type", 18), ("S/N", 18), ("Calibration due", 13), ("Maintenance due", 13), ("State", 28)]
    width = len(headings)
    _title(sheet, "Stations", subtitle, width)
    for column, (_, col_width) in enumerate(headings, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = col_width

    row = 4
    for station in stations:
        band = " · ".join(x for x in [
            station.code, station.customer.name if station.customer else "",
            station.process.name, station.line.name if station.line_id else "",
        ] if x)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
        for column in range(1, width + 1):
            _cell(sheet, row, column, band if column == 1 else None, bold=True, colour="FFFFFF", fill=DARK, size=12)
        sheet.row_dimensions[row].height = 22
        row += 1
        for column, (heading, _) in enumerate(headings, start=1):
            _cell(sheet, row, column, heading, bold=True, colour=DARK, fill="DCE6F1", align="center")
        row += 1

        def machine_row(used_as, asset, fill):
            nonlocal row
            calibration = _first_due(asset.calibration_plans.all())
            maintenance = _first_due(asset.maintenance_plans.all())
            problem = unusable_reason(asset)
            values = [None, used_as, asset.equipment_no, asset.code, asset.name, asset.category.name,
                      asset.serial_number,
                      "Unknown" if asset.calibration_date_unknown else (calibration.next_due_at if calibration else None),
                      maintenance.next_due_at if maintenance else None, problem or "OK"]
            for column, value in enumerate(values, start=1):
                heading = headings[column - 1][0]
                colour, bold, fmt, align = None, False, None, "left"
                if heading.endswith("due"):
                    fmt, align = "dd/mm/yyyy", "center"
                    colour = _due_colour(value, today)
                    bold = colour is not None
                elif heading == "State":
                    colour, bold = (RED, True) if problem else (GREEN, False)
                elif heading == "Equipment No":
                    bold = True
                _cell(sheet, row, column, value, bold=bold, colour=colour, fill=fill, number_format=fmt, align=align)
            _photo(sheet, f"A{row}", asset)
            sheet.row_dimensions[row].height = ROW_PX * 0.75
            row += 1

        for inst in station.open_installations().select_related("component__category").prefetch_related(
            "component__calibration_plans", "component__maintenance_plans", "component__attachments",
        ):
            machine_row("Shared", inst.component, None)

        for test_set in station.test_sets.select_related("product_process__product").order_by(
            "product_process__product__code", "number"
        ):
            readiness = evaluate_set(test_set)
            label = f"{test_set.product.code} #{test_set.number}"
            for line in readiness.lines:
                if line.from_station:
                    continue
                if line.equipment:
                    machine_row(f"{label} · {line.role.name}", line.equipment, ZEBRA)
                else:
                    values = [None, f"{label} · {line.role.name}", "", "", "— nothing fitted —", line.role.name,
                              "", None, None, "Missing" if line.is_blocking else "Optional, empty"]
                    for column, value in enumerate(values, start=1):
                        missing = column == width and line.is_blocking
                        _cell(sheet, row, column, value, fill=ZEBRA, colour=RED if missing else GREY, bold=missing)
                    row += 1
            state = "Ready" if readiness.is_ready else readiness.summary
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=width)
            for column in range(1, width + 1):
                _cell(sheet, row, column, None if column != 2 else f"{label}: {state}", bold=True,
                      colour=GREEN if readiness.is_ready else RED, fill="F2F2F2" if column > 1 else None,
                      border=column > 1)
            row += 1
        row += 1                                   # a blank line between stations
    sheet.freeze_panes = "A4"
    _print_setup(sheet, 3)


def _summary_sheet(sheet, assets, at, subtitle):
    """Per type, all counted by formulas on the Equipment sheet: edit it and these follow."""
    def col(group, heading):
        letter = at[(group, heading)]
        return f"Equipment!${letter}$5:${letter}${LAST_ROW}"

    refs = {
        "kind": col("Description", "Type"), "status": col("Where & state", "Status"),
        "needs": col("Calibration", "Needs"), "cal_due": col("Calibration", "Due"),
        "maint_due": col("Maintenance", "Due"), "can": col("Where & state", "Can be used"),
        "station": col("In use at", "Station"),
    }
    headings = [
        ("Type", 26, None, False),
        ("Assets", 9, "COUNTIF({kind},{a})", False),
        ("In use", 9, 'COUNTIFS({kind},{a},{status},"In use")', False),
        ("Spare", 9, 'COUNTIFS({kind},{a},{status},"Spare")', False),
        ("Broken", 9, 'COUNTIFS({kind},{a},{status},"Broken")', True),
        ("At a station", 10, 'COUNTIFS({kind},{a},{station},"?*")', False),
        ("Needs calibration", 11, 'COUNTIFS({kind},{a},{needs},"Yes")', False),
        ("Calibration overdue", 11, 'COUNTIFS({kind},{a},{needs},"Yes",{cal_due},"<"&TODAY())', True),
        ("Calibration date unknown", 11,
         'COUNTIFS({kind},{a},{needs},"Yes",{cal_due},"Unknown")+COUNTIFS({kind},{a},{needs},"Yes",{cal_due},"")', True),
        ("Calibration due in 30 days", 11,
         'COUNTIFS({kind},{a},{needs},"Yes",{cal_due},">="&TODAY(),{cal_due},"<="&TODAY()+30)', False),
        ("Maintenance overdue", 11, 'COUNTIFS({kind},{a},{maint_due},"<"&TODAY())', True),
        ("Cannot be used", 10, 'COUNTIFS({kind},{a},{can},"<>OK")', True),
    ]
    width = len(headings)
    _title(sheet, "Summary by equipment type", subtitle, width)
    for column, (heading, col_width, _, _) in enumerate(headings, start=1):
        _cell(sheet, 4, column, heading, bold=True, colour="FFFFFF", fill=DARK, align="center")
        sheet.column_dimensions[get_column_letter(column)].width = col_width
    sheet.row_dimensions[4].height = 42

    names = sorted({asset.category.name for asset in assets}, key=str.lower)
    row = 5
    for index, name in enumerate(names):
        fill = ZEBRA if index % 2 else None
        _cell(sheet, row, 1, name, bold=True, fill=fill)
        for column, (_, _, formula, _) in enumerate(headings[1:], start=2):
            _cell(sheet, row, column, "=" + formula.format(a=f"$A{row}", **refs), fill=fill, align="center")
        row += 1
    _cell(sheet, row, 1, "Total", bold=True, fill="DCE6F1")
    for column in range(2, width + 1):
        letter = get_column_letter(column)
        _cell(sheet, row, column, f"=SUM({letter}5:{letter}{row - 1})", bold=True, fill="DCE6F1", align="center")
    # A problem count above zero stands out.
    for column, (_, _, _, alert) in enumerate(headings, start=1):
        if alert:
            letter = get_column_letter(column)
            sheet.conditional_formatting.add(
                f"{letter}5:{letter}{row}",
                FormulaRule(formula=[f"{letter}5>0"], font=Font(color=RED, bold=True),
                            fill=PatternFill("solid", fgColor="FDE2E1")),
            )
    sheet.freeze_panes = "B5"


def _print_setup(sheet, header_row):
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.sheet_view.zoomScale = 90


def build_register(assets, stations=None, filters_text: str = "") -> Workbook:
    """
    The register workbook for ``assets`` (an Equipment queryset) and
    ``stations`` (all of them when None).
    """
    from systems.models import Installation, Station

    today = timezone.localdate()
    open_fittings = Installation.objects.filter(removed_at__isnull=True).select_related(
        "station__customer", "station__process", "slot__category",
        "test_set__station", "test_set__product_process__product__customer", "test_set__product_process__process",
    )
    assets = list(
        assets.select_related("category")
        .prefetch_related(
            "attachments", "calibration_plans", "maintenance_plans",
            Prefetch("installations", queryset=open_fittings, to_attr="open_fittings"),
        )
        .order_by("category__name", "equipment_no", "code")
    )
    if stations is None:
        stations = Station.objects.all()
    stations = list(stations.select_related("customer", "process", "line").order_by("customer__name", "process__position", "code"))
    station_models = {station.pk: station.product_codes() for station in stations}
    for station in Station.objects.exclude(pk__in=station_models):
        station_models[station.pk] = station.product_codes()

    stamp = timezone.localtime().strftime("%d %b %Y %H:%M")
    subtitle = f"{settings.EMA_LABEL_COMPANY} · exported {stamp}" + (f" · {filters_text}" if filters_text else "")

    workbook = Workbook()
    at = _equipment_sheet(workbook.active, assets, station_models, today, f"{len(assets)} assets · {subtitle}")
    workbook.active.title = "Equipment"
    _stations_sheet(workbook.create_sheet("Stations"), stations, today, f"{len(stations)} stations · {subtitle}")
    _summary_sheet(workbook.create_sheet("Summary"), assets, at, subtitle)
    # Formulas are worked out by Excel when the file opens.
    workbook.calculation.fullCalcOnLoad = True
    return workbook
