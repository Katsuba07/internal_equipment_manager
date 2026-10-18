"""Equipment tests: status changes and spreadsheet import/export."""

import io
import shutil
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from PIL import Image as PILImage

from core.models import Category, Location, Role, User
from core.photos import save_photo
from equipment.models import Equipment, EquipmentStatus
from equipment.register import build_register
from equipment.xlsx_images import cell_images, save_workbook

PASSWORD = "test-pass-2026"
EXPORT_TEMP_MEDIA = tempfile.mkdtemp(prefix="ema-test-export-media-")


def photo_file(color="red"):
    buffer = io.BytesIO()
    PILImage.new("RGB", (800, 600), color).save(buffer, format="PNG")
    return SimpleUploadedFile("asset.png", buffer.getvalue(), content_type="image/png")


class StatusChangeTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(
            username="mgr", password=PASSWORD, full_name="Mai Le", role=Role.MANAGER
        )
        cls.equipment = Equipment.objects.create(
            equipment_no="FCT-0001", name="FCT station 1", plant="CTV2",
            category=Category.objects.create(code="TESTER", name="Test system"),
        )

    def test_a_new_asset_is_in_use(self):
        self.assertEqual(self.equipment.status, EquipmentStatus.IN_USE)

    def test_change_status_records_history(self):
        entry = self.equipment.change_status(EquipmentStatus.BROKEN, self.user, reason="Intermittent fail on pin 45")
        self.equipment.refresh_from_db()
        self.assertEqual(self.equipment.status, EquipmentStatus.BROKEN)
        self.assertEqual((entry.from_status, entry.to_status), (EquipmentStatus.IN_USE, EquipmentStatus.BROKEN))
        self.assertEqual(entry.changed_by, self.user)
        self.assertEqual(self.equipment.status_history.count(), 1)

    def test_changing_to_same_status_does_nothing(self):
        self.assertIsNone(self.equipment.change_status(EquipmentStatus.IN_USE, self.user, reason="No change"))
        self.assertEqual(self.equipment.status_history.count(), 0)

    def test_status_is_picked_on_the_form_and_kept_in_the_history(self):
        self.client.login(username="mgr", password=PASSWORD)
        self.client.post(reverse("equipment:update", args=[self.equipment.pk]), {
            "equipment_no": "FCT-0001", "name": "FCT station 1", "category": self.equipment.category.pk,
            "plant": "CTV2", "status": EquipmentStatus.SPARE,
            "parts-TOTAL_FORMS": "0", "parts-INITIAL_FORMS": "0",
        })
        self.equipment.refresh_from_db()
        self.assertEqual(self.equipment.status, EquipmentStatus.SPARE)
        entry = self.equipment.status_history.get()
        self.assertEqual((entry.from_status, entry.to_status), (EquipmentStatus.IN_USE, EquipmentStatus.SPARE))
        body = self.client.get(reverse("equipment:list") + "?status=SPARE").content.decode()
        self.assertIn("FCT-0001", body)

    def test_a_broken_asset_cannot_be_used(self):
        self.assertTrue(self.equipment.is_available)
        self.equipment.change_status(EquipmentStatus.BROKEN, self.user, reason="Fixture fault")
        self.assertFalse(self.equipment.is_available)


class EquipmentFormTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.category = Category.objects.create(code="FIX-ICT", name="ICT fixture")
        cls.location = Location.objects.create(code="ICT-RM", name="ICT Room")
        User.objects.create_user(
            username="manager", password=PASSWORD, full_name="Minh Nguyen", role=Role.MANAGER
        )

    def test_create_equipment_adds_opening_history_entry(self):
        self.client.login(username="manager", password=PASSWORD)
        response = self.client.post(
            reverse("equipment:create"),
            {
                "code": "ict-fix-007", "equipment_no": "ict-fix-007",
                "name": "ICT fixture — Sensor board",
                "category": self.category.pk,
                "location": self.location.pk,
                "year_made": "",
                "plant": "CTV2", "parts-TOTAL_FORMS": "0",
                "parts-INITIAL_FORMS": "0",
            },
        )
        self.assertEqual(response.status_code, 302)
        asset = Equipment.objects.get(code="ICT-FIX-007")  # normalised to upper case
        self.assertEqual(asset.status_history.count(), 1)
        self.assertEqual(asset.status_history.first().reason, "Asset record created")


@override_settings(MEDIA_ROOT=EXPORT_TEMP_MEDIA)
class ExportWithPhotosTests(TestCase):
    """The .xlsx export drops each asset's main photo into its row."""

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(EXPORT_TEMP_MEDIA, ignore_errors=True)

    def setUp(self):
        category = Category.objects.create(code="TESTER", name="Test system")
        self.with_photo = Equipment.objects.create(
            code="FCT-STN-001", name="Has a photo", category=category
        )
        self.without_photo = Equipment.objects.create(
            code="FCT-STN-002", name="No photo yet", category=category
        )
        save_photo(self.with_photo, photo_file(), None)

    HEAD = 4                                  # title, subtitle, group band, then the headings

    def row(self, sheet, number):
        return {sheet.cell(row=self.HEAD, column=c).value: sheet.cell(row=number, column=c)
                for c in range(1, sheet.max_column + 1)}

    def test_register_embeds_the_main_photo_under_grouped_headings(self):
        sheet = build_register(Equipment.objects.all())["Equipment"]
        self.assertEqual(sheet.cell(row=1, column=1).value, "Equipment register")
        self.assertEqual(sheet.cell(row=3, column=1).value, "Asset")
        self.assertEqual([sheet.cell(row=4, column=c).value for c in (1, 2, 3)], ["Photo", "Equipment No", "ZT Code"])
        self.assertEqual(sheet._images, [])                       # nothing floating over the sheet
        self.assertEqual(list(cell_images(sheet.parent)), [("Equipment", "A5")])   # FCT-STN-001, the first asset
        self.assertEqual(sheet.auto_filter.ref.split(":")[0], "A4")

    def test_the_saved_file_places_the_photo_in_its_cell(self):
        import zipfile

        out = io.BytesIO()
        save_workbook(build_register(Equipment.objects.all()), out)
        package = zipfile.ZipFile(out)
        names = package.namelist()
        for part in ("xl/metadata.xml", "xl/richData/rdrichvalue.xml", "xl/richData/richValueRel.xml",
                     "xl/media/cellimage1.jpeg"):
            self.assertIn(part, names)
        sheet_xml = package.read("xl/worksheets/sheet1.xml").decode()
        self.assertRegex(sheet_xml, r'<c r="A5"[^>]* t="e" vm="1"><v>#VALUE!</v></c>')
        self.assertIn("sheetMetadata", package.read("xl/_rels/workbook.xml.rels").decode())
        self.assertIn("rdrichvalue+xml", package.read("[Content_Types].xml").decode())
        book = load_workbook(io.BytesIO(out.getvalue()))           # still a valid workbook
        self.assertEqual(book["Equipment"]["B4"].value, "Equipment No")

    def test_register_without_a_photo_does_not_crash(self):
        book = build_register(Equipment.objects.filter(pk=self.without_photo.pk))
        self.assertEqual(cell_images(book), {})

    def test_register_leaves_out_the_retired_columns(self):
        sheet = build_register(Equipment.objects.filter(pk=self.with_photo.pk))["Equipment"]
        headers = [cell.value for cell in sheet[self.HEAD]]
        for gone in ("Firmware version", "Purchase date", "Purchase cost (USD)", "Warranty until", "Supplier"):
            self.assertNotIn(gone, headers)
        self.assertEqual(headers[-2:], ["Specifications", "Notes"])

    def test_register_carries_calibration_and_maintenance_dates(self):
        from datetime import date

        from calibration.models import CalibrationPlan
        from maintenance.models import MaintenancePlan

        asset = self.with_photo
        CalibrationPlan.objects.create(
            equipment=asset, name="Calibration", interval_months=12,
            last_calibrated_at=date(2025, 1, 10), next_due_at=date(2026, 1, 10),
        )
        MaintenancePlan.objects.create(equipment=asset, name="Clean", interval_days=90, next_due_at=date(2099, 1, 1))
        sheet = build_register(Equipment.objects.filter(pk=asset.pk))["Equipment"]
        cells = [sheet.cell(row=5, column=c) for c in range(1, sheet.max_column + 1)]
        headers = [cell.value for cell in sheet[self.HEAD]]
        due = [cells[i] for i, h in enumerate(headers) if h == "Due"]
        every = [cells[i].value for i, h in enumerate(headers) if h.startswith("Every")]
        self.assertEqual(every, [12, 90])
        self.assertEqual([c.value for c in due], [date(2026, 1, 10), date(2099, 1, 1)])
        # Past dates turn red by a conditional format against TODAY(), not a fixed colour.
        rules = [rule.formula[0] for cf in sheet.conditional_formatting for rule in cf.rules]
        self.assertIn(f"AND(ISNUMBER({due[0].column_letter}5),{due[0].column_letter}5<TODAY())", rules)
        self.assertIn("Calibration", [c.value for c in sheet[3]])
        # Can be used is plain words as of the export day, no day counts.
        self.assertEqual(cells[headers.index("Can be used")].value, "Calibration overdue")

    def test_register_says_where_a_machine_is_in_use(self):
        from core.models import Customer, Product
        from systems.models import Installation, ProductProcess, Station, StationSlot, SystemType, TestSet

        fct = SystemType.objects.create(code="FCT", name="Functional test")
        baf = Customer.objects.create(code="BAF", name="BAF")
        station = Station.objects.create(code="BAF-FCT-01", process=fct, customer=baf)
        station.set_equipment([self.with_photo], None)
        slot = StationSlot.objects.create(station=station, category=self.without_photo.category)
        step = ProductProcess.objects.create(product=Product.objects.create(code="BAF7350", customer=baf), process=fct)
        test_set = TestSet.objects.create(product_process=step, number=1, station=station)
        Installation.fit(self.without_photo, test_set, slot, None)
        book = build_register(Equipment.objects.all())
        sheet = book["Equipment"]
        rows = {sheet.cell(row=r, column=3).value: self.row(sheet, r) for r in range(5, sheet.max_row + 1)}   # by ZT Code
        shared, own = rows["FCT-STN-001"], rows["FCT-STN-002"]
        self.assertEqual((shared["Station"].value, shared["Used as"].value, shared["Models"].value), ("BAF-FCT-01", "Shared", "BAF7350"))
        self.assertEqual((own["Station"].value, own["Customer"].value, own["Models"].value), ("BAF-FCT-01", "BAF", "BAF7350"))
        stations = book["Stations"]
        text = " ".join(str(c.value) for row in stations.iter_rows() for c in row if c.value)
        for expected in ("BAF-FCT-01", "Shared", "BAF7350 #1", "FCT-STN-002"):
            self.assertIn(expected, text)
        self.assertEqual(book["Summary"].cell(row=1, column=1).value, "Summary by equipment type")

    def test_export_button_downloads_the_register(self):
        User.objects.create_user(username="viewer", password=PASSWORD, full_name="Linh Vo", role=Role.VIEWER)
        self.client.login(username="viewer", password=PASSWORD)
        response = self.client.get(reverse("equipment:export"), {"q": "FCT-STN"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("equipment-register-", response["Content-Disposition"])


class ZtCodeAndLabelTests(TestCase):
    """ZT Code is the key, Equipment No sits beside it, and the label carries both."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(
            username="viewer", password=PASSWORD, full_name="Linh Vo", role=Role.VIEWER
        )
        User.objects.create_user(
            username="tech", password=PASSWORD, full_name="Minh Tran", role=Role.TECHNICIAN
        )
        category = Category.objects.create(code="DMM", name="Digital multimeter")
        cls.dmm = Equipment.objects.create(
            code="70-000148", equipment_no="DMM-001", name="KEYSIGHT 34461A",
            category=category, serial_number="MY57224741",
        )
        cls.old = Equipment.objects.create(code="CTV20200077", name="Old style", category=category)

    def setUp(self):
        self.client.login(username="viewer", password=PASSWORD)

    def test_label_shows_the_four_lines_and_the_qr_link(self):
        response = self.client.get(reverse("equipment:label", args=[self.dmm.pk]))
        for text in ("COMPUTIME VIETNAM", "DMM-001", "KEYSIGHT 34461A", "70-000148", "MY57224741", "<svg"):
            self.assertContains(response, text)
        self.assertEqual(
            self.dmm.label_qr_url,
            "https://computime.datafill.com/AssetManagement_VN?Param=70-000148&Task=AssetNo",
        )

    def test_zebra_label_holds_one_copy_per_label_across_the_roll(self):
        from equipment.labels import LabelSize, build_zpl

        zpl = build_zpl(self.dmm, LabelSize(w=45, h=25, cols=2))
        self.assertTrue(zpl.startswith("^XA") and zpl.rstrip().endswith("^XZ"))
        self.assertIn("^PW1063", zpl)              # 90 mm across at 300 dpi
        self.assertEqual(zpl.count("^GFA"), 2)      # the QR code, drawn as a graphic
        for text in ("COMPUTIME VIETNAM", "DMM-001", "KEYSIGHT 34461A", "MY57224741"):
            self.assertEqual(zpl.count(text), 2, text)

    def test_zebra_label_wraps_a_long_name_on_two_lines_and_escapes_printer_commands(self):
        from equipment.labels import LabelSize, build_zpl

        self.dmm.name = "A^B " + "very long name " * 10
        zpl = build_zpl(self.dmm, LabelSize(w=45, h=25, cols=1))
        self.assertIn("A\\5EB", zpl)
        self.assertNotIn("A^B", zpl)
        self.assertIn(",2,0,L^FH", zpl)                       # a two-line field block
        self.assertNotIn("very long name " * 6, zpl)          # cut to what two lines hold
        short = build_zpl(Equipment(name="KEYSIGHT 34461A", equipment_no="DMM-001"), LabelSize(w=45, h=25, cols=1))
        self.assertNotIn(",2,0,L^FH", short)                   # a short name stays on one line

    def test_a_list_of_assets_fills_the_labels_across_then_the_next_row(self):
        from equipment.labels import LabelSize, build_zpl

        assets = [Equipment(name=f"Meter {n}", equipment_no=f"DMM-00{n}") for n in range(1, 4)]
        zpl = build_zpl(assets, LabelSize(w=45, h=25, cols=2))
        self.assertEqual(zpl.count("^XA"), 2)                  # 3 labels, 2 across: two rows
        first, second = zpl.split("^XZ")[:2]
        self.assertIn("DMM-001", first)
        self.assertIn("DMM-002", first)
        self.assertIn("DMM-003", second)

    def test_zebra_label_text_size_can_be_chosen(self):
        from equipment.labels import LabelSize, build_zpl

        automatic = build_zpl(self.dmm, LabelSize.from_values({"w": "45", "h": "25", "font": "0"}))
        small = build_zpl(self.dmm, LabelSize.from_values({"w": "45", "h": "25", "font": "2"}))
        self.assertIn("^A0N,29,29^FDEQ No:", automatic)
        self.assertIn("^A0N,24,24^FDEQ No:", small)     # 2 mm at 300 dpi

    def test_label_zpl_is_served_for_a_printer_on_the_users_computer(self):
        response = self.client.get(
            reverse("equipment:label_zpl", args=[self.dmm.pk]) + "?w=45&h=25&cols=1&font=2"
        )
        body = response.content.decode()
        self.assertTrue(body.startswith("^XA"))
        self.assertIn("^PW531", body)
        self.assertIn("KEYSIGHT 34461A", body)

    def test_label_page_offers_the_server_printer_only_when_one_is_set(self):
        url = reverse("equipment:label", args=[self.dmm.pk])
        self.client.login(username="tech", password=PASSWORD)
        self.assertNotContains(self.client.get(url), 'id="zebra-print"')
        with self.settings(EMA_LABEL_PRINTER="ZDesigner ZT510-300dpi ZPL"):
            self.assertContains(self.client.get(url), 'id="zebra-print"')
        self.assertContains(self.client.get(url), 'id="local-print"')

    def test_print_label_sends_the_posted_size_to_the_printer(self):
        from unittest import mock

        self.client.login(username="tech", password=PASSWORD)
        with mock.patch("equipment.views.send_to_printer") as send:
            response = self.client.post(
                reverse("equipment:label_print", args=[self.dmm.pk]),
                {"w": "50", "h": "25", "cols": "1", "gap": "0", "dx": "nonsense"},
            )
        self.assertEqual(response.json()["ok"], True)
        self.assertIn("^PW591", send.call_args.args[0])   # one 50 mm label

    def test_print_label_reports_a_printer_that_cannot_be_reached(self):
        from unittest import mock

        from equipment.labels import LabelPrintError

        self.client.login(username="tech", password=PASSWORD)
        with mock.patch("equipment.views.send_to_printer", side_effect=LabelPrintError("Printer is off.")):
            response = self.client.post(reverse("equipment:label_print", args=[self.dmm.pk]))
        self.assertEqual(response.status_code, 502)
        self.assertEqual(response.json()["message"], "Printer is off.")

    def test_search_finds_an_asset_by_its_equipment_no(self):
        body = self.client.get(reverse("equipment:list") + "?q=dmm-001").content.decode()
        self.assertIn("70-000148", body)
        self.assertNotIn("CTV20200077", body)

    def test_the_list_shows_the_equipment_no_while_the_zt_code_is_blank(self):
        Equipment.objects.create(equipment_no="OSC-0001", name="UNI-T UPO2102CS", category=self.dmm.category)
        body = self.client.get(reverse("equipment:list")).content.decode()
        self.assertIn('class="mono cell-id"', body)
        self.assertIn("OSC-0001", body)

    def test_two_assets_cannot_share_an_equipment_no_but_blank_is_fine(self):
        from django.core.exceptions import ValidationError

        clash = Equipment(code="70-000149", equipment_no="DMM-001", name="x", category=self.dmm.category)
        with self.assertRaises(ValidationError):
            clash.full_clean()
        Equipment.objects.create(code="70-000150", name="No number yet", category=self.dmm.category)
        self.assertEqual(Equipment.objects.filter(equipment_no="").count(), 2)

    def test_a_retired_ctv_id_is_flagged_as_pending(self):
        self.assertTrue(self.old.zt_code_pending)
        self.assertFalse(self.dmm.zt_code_pending)
        self.assertContains(self.client.get(self.old.get_absolute_url()), "ZT Code pending")
        self.assertNotContains(self.client.get(self.dmm.get_absolute_url()), "ZT Code pending")


class HeaderFilterTests(TestCase):
    """The list's filters sit in the column headers and still filter."""

    @classmethod
    def setUpTestData(cls):
        User.objects.create_user(username="viewer", password=PASSWORD, full_name="Linh Vo", role=Role.VIEWER)
        cls.dmm = Category.objects.create(code="DMM", name="Digital multimeter")
        psu = Category.objects.create(code="PSU", name="Power supply")
        Equipment.objects.create(code="70-000001", name="Meter", category=cls.dmm)
        Equipment.objects.create(code="70-000002", name="Supply", category=psu)

    def setUp(self):
        self.client.login(username="viewer", password=PASSWORD)

    def test_each_filter_is_in_a_header_and_tied_to_the_search_form(self):
        body = self.client.get(reverse("equipment:list")).content.decode()
        self.assertIn('<form method="get" class="toolbar no-print" id="list-filter" data-autosubmit>', body)
        head = body[body.index("<thead>"):body.index("</thead>")]
        for name in ("category", "plant", "status", "labelled", "serial", "is_measuring_device"):
            self.assertRegex(head, rf'<select name="{name}"[^>]*form="list-filter"')
        self.assertIn('<option value="" selected>Category</option>', head)
        toolbar = body[body.index('id="list-filter"'):body.index('id="results"')]
        self.assertNotIn("<select", toolbar)

    def test_choosing_a_header_filter_narrows_the_list_and_keeps_its_value(self):
        body = self.client.get(reverse("equipment:list"), {"category": self.dmm.pk}).content.decode()
        self.assertIn("70-000001", body)
        self.assertNotIn("70-000002", body)
        self.assertIn(f'<option value="{self.dmm.pk}" selected>', body)


class EquipmentNoCounterTests(TestCase):
    """Equipment No runs per type; the form shows the last one and offers the next."""

    @classmethod
    def setUpTestData(cls):
        User.objects.create_user(username="viewer", password=PASSWORD, full_name="Linh Vo", role=Role.VIEWER)
        cls.dmm = Category.objects.create(code="DMM", name="Digital multimeter")
        cls.psu = Category.objects.create(code="PSU", name="Power supply")
        for code, number in (("70-1", "DMM-001"), ("70-2", "DMM-009"), ("70-3", "DMM-X"), ("70-4", "")):
            Equipment.objects.create(code=code, name=code, category=cls.dmm, equipment_no=number)

    def test_next_follows_the_highest_number_of_the_type(self):
        counter = Equipment.equipment_no_counter(self.dmm)
        self.assertEqual((counter["last"], counter["next"], counter["count"]), ("DMM-0009", "DMM-0010", 2))
        self.assertEqual(Equipment.equipment_no_counter(self.psu)["next"], "PSU-0001")

    def test_numbers_are_padded_to_four_digits(self):
        for typed, kept in (("dmm-1", "DMM-0001"), (" DMM-012 ", "DMM-0012"), ("DMM-12345", "DMM-12345"),
                            ("FCT-FIX-9", "FCT-FIX-9"), ("NPI RHEEM", "NPI RHEEM"), ("", "")):
            self.assertEqual(Equipment.normalise_equipment_no(typed), kept)

    def test_the_form_asks_for_it_by_type(self):
        self.client.login(username="viewer", password=PASSWORD)
        data = self.client.get(reverse("equipment:next_no"), {"category": self.dmm.pk}).json()
        self.assertEqual(data["next"], "DMM-0010")
        self.assertEqual(self.client.get(reverse("equipment:next_no")).json()["next"], "")


class LabelPopupTests(TestCase):
    def test_print_label_opens_in_a_pop_up_that_may_frame_the_label(self):
        User.objects.create_user(username="viewer", password=PASSWORD, full_name="Linh Vo", role=Role.VIEWER)
        asset = Equipment.objects.create(code="70-000148", name="DMM", category=Category.objects.create(code="DMM", name="DMM"))
        self.client.login(username="viewer", password=PASSWORD)
        page = self.client.get(asset.get_absolute_url()).content.decode()
        self.assertIn('data-frame-dialog="Label 70-000148"', page)
        self.assertNotIn('target="_blank"', page.split("Print label")[0][-400:])
        label = self.client.get(reverse("equipment:label", args=[asset.pk]))
        self.assertEqual(label["X-Frame-Options"], "SAMEORIGIN")
        self.assertEqual(self.client.get(asset.get_absolute_url())["X-Frame-Options"], "DENY")   # other pages stay unframeable


class CalibrationByTypeTests(TestCase):
    """A type marked as needing calibration starts its assets with the tick on."""

    @classmethod
    def setUpTestData(cls):
        User.objects.create_user(username="mgr", password=PASSWORD, full_name="Mai Le", role=Role.MANAGER)
        cls.dmm = Category.objects.create(code="DMM", name="Digital multimeter", needs_calibration=True)
        cls.fx = Category.objects.create(code="FX", name="Fixture")
        cls.user = User.objects.get(username="mgr")

    def test_the_form_knows_which_types_need_it(self):
        self.client.login(username="mgr", password=PASSWORD)
        page = self.client.get(reverse("equipment:create")).content.decode()
        self.assertIn(f'data-calibration-types="{self.dmm.pk}"', page)

    def post(self, **fields):
        data = {"equipment_no": "DMM-0001", "category": self.dmm.pk, "name": "METER",
                "plant": "CTV2", "parts-TOTAL_FORMS": "0", "parts-INITIAL_FORMS": "0", "parts-MIN_NUM_FORMS": "0", "parts-MAX_NUM_FORMS": "1000"}
        data.update(fields)
        self.client.login(username="mgr", password=PASSWORD)
        return self.client.post(reverse("equipment:create"), data)

    def test_the_server_will_not_save_an_unticked_instrument_without_a_yes(self):
        page = self.post().content.decode()                      # DMM, calibration not ticked
        self.assertFalse(Equipment.objects.exists())
        self.assertIn("normally needs calibration", page)
        self.assertIn('name="calibration_confirmed"', page)
        self.post(calibration_confirmed="on")
        self.assertFalse(Equipment.objects.get().is_measuring_device)

    def test_the_server_will_not_save_missing_dates_without_a_yes(self):
        page = self.post(is_measuring_device="on").content.decode()
        self.assertFalse(Equipment.objects.exists())
        self.assertIn("No calibration dates are filled in", page)
        self.post(is_measuring_device="on", calibration_confirmed="on")
        self.assertTrue(Equipment.objects.get().calibration_date_unknown)

    def test_with_dates_nothing_is_asked(self):
        self.post(is_measuring_device="on", calibration_interval_months="12", calibration_last="2026-01-10")
        self.assertFalse(Equipment.objects.get().calibration_date_unknown)

    def test_quick_add_follows_the_type(self):
        from core.quickadd import EquipmentQuickForm

        for category, expected in ((self.dmm, True), (self.fx, False)):
            form = EquipmentQuickForm({"code": f"Q-{category.code}", "name": "Quick", "category": category.pk})
            form.user = self.user
            self.assertTrue(form.is_valid(), form.errors)
            self.assertEqual(form.save().is_measuring_device, expected)



class CalibrationDateUnknownTests(TestCase):
    """Needs calibration but no due date on record: treated as not to be used."""

    @classmethod
    def setUpTestData(cls):
        from datetime import date

        from calibration.models import CalibrationPlan

        User.objects.create_user(username="tech", password=PASSWORD, full_name="Khanh Pham", role=Role.TECHNICIAN)
        dmm = Category.objects.create(code="DMM", name="Digital multimeter")
        cls.unknown = Equipment.objects.create(equipment_no="DMM-0001", name="LOST STICKER", category=dmm,
                                               is_measuring_device=True)
        cls.dated = Equipment.objects.create(equipment_no="DMM-0002", name="HAS A DATE", category=dmm,
                                             is_measuring_device=True)
        CalibrationPlan.objects.create(equipment=cls.dated, name="Calibration", interval_months=12,
                                       last_calibrated_at=date(2026, 1, 1), next_due_at=date(2099, 1, 1))
        cls.plain = Equipment.objects.create(equipment_no="FX-0001", name="FIXTURE", category=dmm)

    def test_only_the_one_without_a_date_is_unknown(self):
        self.assertTrue(self.unknown.calibration_date_unknown)
        self.assertFalse(self.dated.calibration_date_unknown)
        self.assertFalse(self.plain.calibration_date_unknown)
        self.assertEqual(list(Equipment.objects.calibration_date_unknown()), [self.unknown])

    def test_it_cannot_be_used(self):
        from systems.services import unusable_reason

        self.assertEqual(unusable_reason(self.unknown), "Calibration date unknown")
        self.assertEqual(unusable_reason(self.dated), "")

    def test_list_marks_it_and_filters_it(self):
        self.client.login(username="tech", password=PASSWORD)
        body = self.client.get(reverse("equipment:list")).content.decode()
        self.assertIn("Date unknown", body)
        body = self.client.get(reverse("equipment:list") + "?is_measuring_device=unknown").content.decode()
        self.assertIn("DMM-0001", body)
        self.assertNotIn("DMM-0002", body)
        body = self.client.get(reverse("equipment:list") + "?is_measuring_device=true").content.decode()
        self.assertIn("DMM-0002", body)

    def test_list_shows_where_each_calibration_stands(self):
        from datetime import timedelta

        from django.utils import timezone

        from calibration.models import CalibrationPlan

        late = Equipment.objects.create(equipment_no="DMM-0003", name="LATE", category=self.dated.category,
                                        is_measuring_device=True)
        CalibrationPlan.objects.create(equipment=late, name="Calibration", interval_months=12,
                                       last_calibrated_at=timezone.localdate() - timedelta(days=400),
                                       next_due_at=timezone.localdate() - timedelta(days=12))
        soon = Equipment.objects.create(equipment_no="DMM-0004", name="SOON", category=self.dated.category,
                                        is_measuring_device=True)
        CalibrationPlan.objects.create(equipment=soon, name="Calibration", interval_months=12,
                                       last_calibrated_at=timezone.localdate() - timedelta(days=340),
                                       next_due_at=timezone.localdate() + timedelta(days=20))
        self.client.login(username="tech", password=PASSWORD)
        body = self.client.get(reverse("equipment:list")).content.decode()
        self.assertIn('<span class="badge badge-danger" title="Calibration due', body)
        self.assertIn("12 days overdue", body)
        self.assertIn('<span class="badge badge-warning" title="Calibration due', body)     # within a month
        self.assertIn("In 20 days", body)
        self.assertIn('title="Calibration due 01/01/2099"', body)
        self.assertNotIn(">Required<", body)

    def test_detail_warns_and_dashboard_counts_it(self):
        self.client.login(username="tech", password=PASSWORD)
        self.assertContains(self.client.get(self.unknown.get_absolute_url()), "Calibration date unknown")
        self.assertNotContains(self.client.get(self.dated.get_absolute_url()), "Calibration date unknown")
        self.assertEqual(self.client.get(reverse("dashboard")).context["calibration_unknown_count"], 1)

    def test_export_says_unknown(self):
        sheet = build_register(Equipment.objects.all())["Equipment"]
        heads = [c.value for c in sheet[4]]
        due = [i for i, h in enumerate(heads, start=1) if h == "Due"][0]       # the calibration Due
        rows = {sheet.cell(r, 2).value: sheet.cell(r, due) for r in range(5, sheet.max_row + 1)}
        self.assertEqual(rows["DMM-0001"].value, "Unknown")
        last = [i for i, h in enumerate(heads, start=1) if h == "Last"][0]       # the calibration Last
        self.assertEqual(sheet.cell(rows["DMM-0001"].row, last).value, "Unknown")
        self.assertNotEqual(sheet.cell(rows["DMM-0002"].row, last).value, "Unknown")

    def test_summary_is_all_formulas(self):
        book = build_register(Equipment.objects.all())
        summary = book["Summary"]
        heads = [c.value for c in summary[4]]
        row = next(r for r in range(5, summary.max_row + 1) if summary.cell(r, 1).value == "Digital multimeter")
        for column in range(2, len(heads) + 1):
            self.assertTrue(str(summary.cell(row, column).value).startswith("="), heads[column - 1])
        self.assertIn("Equipment!", summary.cell(row, 2).value)
        self.assertTrue(book.calculation.fullCalcOnLoad)


class LabelledTests(TestCase):
    """The list's Label tick marks whether the asset has its label on."""

    @classmethod
    def setUpTestData(cls):
        User.objects.create_user(username="tech", password=PASSWORD, full_name="Khanh Pham", role=Role.TECHNICIAN)
        User.objects.create_user(username="viewer", password=PASSWORD, full_name="Linh Vo", role=Role.VIEWER)
        category = Category.objects.create(code="DMM", name="Digital multimeter")
        cls.asset = Equipment.objects.create(equipment_no="DMM-0001", name="KEYSIGHT 34461A", category=category)

    def test_a_technician_ticks_and_unticks_it(self):
        self.client.login(username="tech", password=PASSWORD)
        url = reverse("equipment:labelled", args=[self.asset.pk])
        response = self.client.post(url, {"labelled": "1"})
        self.assertContains(response, "checked")
        self.asset.refresh_from_db()
        self.assertTrue(self.asset.labelled)
        self.client.post(url, {})                          # an unticked box sends nothing
        self.asset.refresh_from_db()
        self.assertFalse(self.asset.labelled)

    def test_the_list_offers_the_tick_and_filters_by_it(self):
        self.client.login(username="tech", password=PASSWORD)
        body = self.client.get(reverse("equipment:list")).content.decode()
        self.assertIn(reverse("equipment:labelled", args=[self.asset.pk]), body)
        self.assertIn("DMM-0001", self.client.get(reverse("equipment:list") + "?labelled=false").content.decode())
        self.assertNotIn("DMM-0001", self.client.get(reverse("equipment:list") + "?labelled=true").content.decode())

    def test_a_printed_label_ticks_the_assets(self):
        other = Equipment.objects.create(equipment_no="DMM-0002", name="Second", category=self.asset.category)
        self.client.login(username="tech", password=PASSWORD)
        page = self.client.get(reverse("equipment:labels") + f"?ids={self.asset.pk},{other.pk}").content.decode()
        self.assertIn(f'data-ids="{self.asset.pk},{other.pk}"', page)
        response = self.client.post(reverse("equipment:mark_labelled"), {"ids": f"{self.asset.pk},{other.pk}"})
        self.assertEqual(sorted(response.json()["ids"]), sorted([self.asset.pk, other.pk]))
        self.assertEqual(Equipment.objects.filter(labelled=True).count(), 2)

    def test_ids_past_a_thousand_reach_the_page_as_plain_digits(self):
        # "1,500" (a thousands separator) silently ticked nothing after a print.
        big = Equipment.objects.create(pk=1500, equipment_no="DMM-1500", name="Big id", category=self.asset.category)
        self.client.login(username="tech", password=PASSWORD)
        page = self.client.get(reverse("equipment:label", args=[big.pk])).content.decode()
        self.assertIn('data-ids="1500"', page)
        self.client.post(reverse("equipment:mark_labelled"), {"ids": "1500"})
        big.refresh_from_db()
        self.assertTrue(big.labelled)

    def test_printing_on_the_server_printer_ticks_it(self):
        from unittest import mock

        self.client.login(username="tech", password=PASSWORD)
        with mock.patch("equipment.views.send_to_printer"):
            response = self.client.post(reverse("equipment:label_print", args=[self.asset.pk]))
        self.assertEqual(response.json()["ids"], [self.asset.pk])
        self.asset.refresh_from_db()
        self.assertTrue(self.asset.labelled)

    def test_a_viewer_sees_it_but_cannot_change_it(self):
        self.client.login(username="viewer", password=PASSWORD)
        body = self.client.get(reverse("equipment:list")).content.decode()
        self.assertIn(f'class="label-tick" data-asset="{self.asset.pk}" disabled', body)
        self.assertNotIn(reverse("equipment:labelled", args=[self.asset.pk]), body)
        response = self.client.post(reverse("equipment:labelled", args=[self.asset.pk]), {"labelled": "1"})
        self.assertIn(response.status_code, (302, 403))
        response = self.client.post(reverse("equipment:mark_labelled"), {"ids": str(self.asset.pk)})
        self.assertIn(response.status_code, (302, 403))
        self.asset.refresh_from_db()
        self.assertFalse(self.asset.labelled)
        label_page = self.client.get(reverse("equipment:label", args=[self.asset.pk])).content.decode()
        self.assertNotIn("data-mark-url", label_page)


class AutoNameTests(TestCase):
    """An asset's name follows Manufacturer + Model unless one is typed."""

    @classmethod
    def setUpTestData(cls):
        User.objects.create_user(username="mgr", password=PASSWORD, full_name="Mai Le", role=Role.MANAGER)
        cls.dmm = Category.objects.create(code="DMM", name="Digital multimeter")

    def post(self, **fields):
        data = {"equipment_no": "DMM-050", "category": self.dmm.pk, "status": "IN_USE",
                "plant": "CTV2", "parts-TOTAL_FORMS": "0", "parts-INITIAL_FORMS": "0", "parts-MIN_NUM_FORMS": "0", "parts-MAX_NUM_FORMS": "1000"}
        data.update(fields)
        self.client.login(username="mgr", password=PASSWORD)
        return self.client.post(reverse("equipment:create"), data)

    def test_an_empty_name_is_made_from_manufacturer_and_model(self):
        self.post(manufacturer="KEYSIGHT", model_number="34461A")
        self.assertEqual(Equipment.objects.get(equipment_no="DMM-0050").name, "KEYSIGHT 34461A")

    def test_status_is_saved_and_filtered(self):
        self.post(name="Old meter", status="SPARE")
        self.assertEqual(Equipment.objects.get(equipment_no="DMM-0050").status, "SPARE")
        body = self.client.get(reverse("equipment:list") + "?status=SPARE").content.decode()
        self.assertIn("DMM-0050", body)
        self.assertNotIn("DMM-0050", self.client.get(reverse("equipment:list") + "?status=IN_USE").content.decode())

    def test_the_plant_is_required(self):
        page = self.post(name="No plant", plant="").content.decode()
        self.assertIn("Choose the plant", page)
        self.assertFalse(Equipment.objects.exists())

    def test_a_typed_name_is_kept(self):
        self.post(name="Bench DMM line 2", manufacturer="KEYSIGHT", model_number="34461A")
        self.assertEqual(Equipment.objects.get(equipment_no="DMM-0050").name, "BENCH DMM LINE 2")

    def test_no_name_and_nothing_to_make_one_from_is_refused(self):
        response = self.post()
        self.assertFalse(Equipment.objects.filter(equipment_no="DMM-0050").exists())
        self.assertContains(response, "Manufacturer and Model")

    def test_equipment_no_is_required(self):
        response = self.post(equipment_no="", manufacturer="FLUKE", model_number="17B+")
        self.assertFalse(Equipment.objects.exists())
        self.assertContains(response, "Equipment No")


class BulkLabelTests(TestCase):
    """Tick assets on the list and print their labels together."""

    @classmethod
    def setUpTestData(cls):
        User.objects.create_user(username="viewer", password=PASSWORD, full_name="Linh Vo", role=Role.VIEWER)
        dmm = Category.objects.create(code="DMM", name="Digital multimeter")
        cls.assets = [
            Equipment.objects.create(equipment_no=f"DMM-00{n}", name=f"Meter {n}", category=dmm) for n in range(1, 4)
        ]

    def setUp(self):
        self.client.login(username="viewer", password=PASSWORD)

    def ids(self):
        return ",".join(str(a.pk) for a in self.assets)

    def test_everyone_can_tick_assets_and_print_their_labels(self):
        page = self.client.get(reverse("equipment:list")).content.decode()
        self.assertIn("data-bulk-labels", page)
        self.assertIn("data-bulk-all", page)
        self.assertNotIn("data-bulk-delete", page)                    # deleting stays with administrators

    def test_the_labels_page_shows_every_ticked_asset(self):
        response = self.client.get(reverse("equipment:labels"), {"ids": self.ids()})
        for asset in self.assets:
            self.assertContains(response, asset.equipment_no)
        self.assertContains(response, "Print 3 labels")
        self.assertEqual(response["X-Frame-Options"], "SAMEORIGIN")
        self.assertEqual(self.client.get(reverse("equipment:labels")).status_code, 400)

    def test_the_labels_go_to_the_printer_two_across(self):
        zpl = self.client.get(reverse("equipment:labels_zpl"), {"ids": self.ids(), "w": "45", "h": "25", "cols": "2"}).content.decode()
        self.assertEqual(zpl.count("^XA"), 2)
        self.assertTrue(all(a.equipment_no in zpl for a in self.assets))


class ScheduleDatesShownTests(TestCase):
    """The detail page and the preview give the last and due dates of calibration and maintenance."""

    @classmethod
    def setUpTestData(cls):
        from datetime import date

        from calibration.models import CalibrationPlan
        from maintenance.models import MaintenancePlan

        User.objects.create_user(username="viewer", password=PASSWORD, full_name="Linh Vo", role=Role.VIEWER)
        kind = Category.objects.create(code="DMM", name="Digital multimeter")
        cls.asset = Equipment.objects.create(equipment_no="DMM-0001", name="METER", category=kind, plant="CTV2",
                                             is_measuring_device=True, maintenance_required=True)
        CalibrationPlan.objects.create(equipment=cls.asset, name="Calibration", interval_months=12,
                                       last_calibrated_at=date(2026, 3, 26), next_due_at=date(2027, 3, 26))
        MaintenancePlan.objects.create(equipment=cls.asset, name="Clean", interval_days=90,
                                       last_done_at=date(2026, 9, 1), next_due_at=date(2026, 11, 30))
        cls.unknown = Equipment.objects.create(equipment_no="DMM-0002", name="LOST STICKER", category=kind,
                                               plant="CTV2", is_measuring_device=True)

    def test_detail_and_preview_show_the_dates(self):
        self.client.login(username="viewer", password=PASSWORD)
        for url in (self.asset.get_absolute_url(), reverse("equipment:preview", args=[self.asset.pk])):
            body = self.client.get(url, HTTP_HX_REQUEST="true").content.decode()
            for text in ("Last <b>26 Mar 2026</b>", "Due <b>26 Mar 2027</b>",
                         "Last <b>01 Sep 2026</b>", "Due <b>30 Nov 2026</b>", "every 90 days"):
                self.assertIn(text, body, url)

    def test_an_unknown_calibration_says_so(self):
        self.client.login(username="viewer", password=PASSWORD)
        body = self.client.get(reverse("equipment:preview", args=[self.unknown.pk]), HTTP_HX_REQUEST="true").content.decode()
        self.assertIn("Last <b>Unknown</b> · Due <b>Unknown</b>", body)


class SerialNumberTests(TestCase):
    """A serial number already on another asset is refused on the form."""

    @classmethod
    def setUpTestData(cls):
        User.objects.create_user(username="mgr", password=PASSWORD, full_name="Mai Le", role=Role.MANAGER)
        cls.kind = Category.objects.create(code="MON", name="Monitor")
        cls.first = Equipment.objects.create(equipment_no="MON-0001", name="ACER V196WL", category=cls.kind,
                                             plant="CTV2", serial_number="MMLY7SS001")

    def post(self, number, serial, pk=None):
        self.client.login(username="mgr", password=PASSWORD)
        url = reverse("equipment:update", args=[pk]) if pk else reverse("equipment:create")
        return self.client.post(url, {
            "equipment_no": number, "name": "ACER V196WL", "category": self.kind.pk, "plant": "CTV2",
            "serial_number": serial, "parts-TOTAL_FORMS": "0", "parts-INITIAL_FORMS": "0",
        })

    def test_a_repeated_serial_is_refused_and_names_the_other_asset(self):
        page = self.post("MON-0002", " mmly7ss001 ").content.decode()
        self.assertIn("S/N MMLY7SS001 is already on MON-0001", page)
        self.assertFalse(Equipment.objects.filter(equipment_no="MON-0002").exists())

    def test_stand_ins_and_new_serials_are_fine(self):
        self.post("MON-0002", "NA")
        self.post("MON-0003", "n/a")
        self.post("MON-0004", "MMLY7SS002")
        self.assertEqual(Equipment.objects.count(), 4)

    def test_an_asset_keeping_its_own_repeated_serial_can_still_be_edited(self):
        twin = Equipment.objects.create(equipment_no="MON-0002", name="ACER V196WL", category=self.kind,
                                        plant="CTV2", serial_number="MMLY7SS001")      # on record from before
        response = self.post("MON-0002", "MMLY7SS001", pk=twin.pk)
        self.assertEqual(response.status_code, 302)

    def test_the_list_finds_repeated_serials(self):
        Equipment.objects.create(equipment_no="MON-0002", name="TWIN", category=self.kind, plant="CTV2",
                                 serial_number="MMLY7SS001")
        Equipment.objects.create(equipment_no="MON-0003", name="ALONE", category=self.kind, plant="CTV2",
                                 serial_number="OTHER")
        self.client.login(username="mgr", password=PASSWORD)
        body = self.client.get(reverse("equipment:list") + "?serial=repeated").content.decode()
        self.assertIn("MON-0001", body)
        self.assertIn("MON-0002", body)
        self.assertNotIn("MON-0003", body)


class LabelSerialRowTests(TestCase):
    """On the label the S/N starts right after its own caption, so a long one keeps its room."""

    def test_serial_starts_before_the_aligned_values(self):
        import re

        from equipment.labels import LabelSize, build_zpl

        serial = "SN-2024-0142-ABCDEFGH"
        zpl = build_zpl(Equipment(equipment_no="DMM-0001", name="KEYSIGHT 34461A", code="70-000148",
                                  serial_number=serial), LabelSize(w=45, h=25, cols=1))
        # Each printed value: "^FO<x>,<y>^A0N,..^FH\^FD<text>^FS".
        x_of = {line.split("^FD")[1].removesuffix("^FS"): int(re.match(r"\^FO(\d+),", line).group(1))
                for line in zpl.splitlines() if "^FD" in line and line.startswith("^FO")}
        self.assertLess(x_of[serial], x_of["DMM-0001"])          # not lined up under the other values
        self.assertIn(serial, zpl)                               # printed whole, not cut short
