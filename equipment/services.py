"""
Spreadsheet import and export for equipment.

Kept out of the views so the parsing can be tested on its own, and reused by
a management command when a bulk load is needed.
"""

from dataclasses import dataclass, field
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

from core.models import Category, Location

from .models import Equipment, EquipmentStatus

# Size, in pixels, of the photo dropped into each export row's first cell.
PHOTO_CELL_PX = 64

# Spreadsheet column → model field. The first three are required.
COLUMNS = [
    ("ZT Code", "code", True),
    ("Name", "name", True),
    ("Category code", "category", True),
    ("Location code", "location", False),
    ("Manufacturer", "manufacturer", False),
    ("Model", "model_number", False),
    ("Serial number", "serial_number", False),
    ("Year", "year_made", False),
    ("Status", "status", False),
    ("Needs calibration (x)", "is_measuring_device", False),
    ("Specifications", "specifications", False),
    ("Notes", "notes", False),
    ("Asset No.", "asset_no", False),
    ("Equipment No", "equipment_no", False),
]

# Older templates called the first column "Asset ID" or "Equipment ID"; they still load.
HEADER_ALIASES = {"asset id": "zt code", "equipment id": "zt code"}

HEADERS = [label for label, _, _ in COLUMNS]

# Accept either the stored code or the human label when reading a status.
STATUS_LOOKUP = {label.lower(): value for value, label in EquipmentStatus.choices}
STATUS_LOOKUP.update({value.lower(): value for value, _ in EquipmentStatus.choices})


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)

    @property
    def has_errors(self) -> bool:
        return bool(self.errors)

    @property
    def total_processed(self) -> int:
        return self.created + self.updated


def _clean(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return value


def _parse_int(value, label: str, row_number: int, errors: list[str]):
    value = _clean(value)
    if value in ("", None):
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        errors.append(f"Row {row_number}: '{label}' is not a whole number.")
        return None


def _is_truthy(value) -> bool:
    return str(_clean(value)).lower() in {"x", "1", "true", "yes", "y"}


def build_template_workbook() -> Workbook:
    """Build the import template, including the catalog codes already defined."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Equipment"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F6FEB")
    required_fill = PatternFill("solid", fgColor="C93C37")

    for index, (label, _, required) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.font = header_font
        cell.fill = required_fill if required else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = max(len(label) + 4, 16)
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A2"

    sheet.append([
        "FCT-FIX-014", "FCT fixture — BAF1950", "FIX-FCT", "TEST-FCT",
        "In-house", "FCT-014", "SN-2024-0142", 2024, "In production",
        "", "128 test points, 4 UUT positions", "Primary fixture for line 2",
    ])

    guide = workbook.create_sheet("Reference")
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 58
    guide.append(["ITEM", "ACCEPTED VALUES"])
    guide["A1"].font = Font(bold=True)
    guide["B1"].font = Font(bold=True)
    guide.append(["Red header columns", "Required — must be filled in"])
    guide.append(["Needs calibration", "Put an x for instruments that require calibration"])
    guide.append([])
    guide.append(["Valid status values", ""])
    for _, label in EquipmentStatus.choices:
        guide.append(["", label])

    for title, queryset in (
        ("Category codes", Category.objects.active()),
        ("Location codes", Location.objects.active()),
    ):
        guide.append([])
        guide.append([title, ""])
        for code, name in queryset.values_list("code", "name"):
            guide.append(["", f"{code} — {name}"])

    return workbook


def import_equipment_from_xlsx(file_obj, user, update_existing: bool = False) -> ImportResult:
    """
    Read a spreadsheet and create or update equipment.

    The whole file runs inside one database transaction: if any row fails,
    nothing is written. People fix the file and upload again rather than
    cleaning up a half-loaded import.
    """
    result = ImportResult()

    try:
        workbook = load_workbook(file_obj, data_only=True, read_only=True)
    except Exception:
        result.errors.append("Could not read the file. Make sure it is a valid .xlsx workbook.")
        return result

    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        result.errors.append("The file is empty.")
        return result

    headers = [str(_clean(cell)).lower() for cell in header_row]
    headers = [HEADER_ALIASES.get(header, header) for header in headers]
    missing = [label for label, _, required in COLUMNS if required and label.lower() not in headers]
    if missing:
        result.errors.append("Missing required columns: " + ", ".join(missing) + ".")
        return result

    index_of = {label.lower(): position for position, label in enumerate(headers)}

    def cell(row, label):
        position = index_of.get(label.lower())
        if position is None or position >= len(row):
            return None
        return row[position]

    # Load catalogs up front instead of querying per row.
    categories = {c.code.upper(): c for c in Category.objects.all()}
    locations = {loc.code.upper(): loc for loc in Location.objects.all()}

    pending_new: list[Equipment] = []
    pending_update: list[Equipment] = []
    seen_codes: set[str] = set()
    seen_asset_nos: set[str] = set()
    seen_equipment_nos: set[str] = set()

    for row_number, row in enumerate(rows, start=2):
        if row is None or all(_clean(value) in ("", None) for value in row):
            continue

        code = str(_clean(cell(row, "ZT Code"))).upper()
        name = str(_clean(cell(row, "Name")))
        if not code or not name:
            result.errors.append(f"Row {row_number}: ZT Code and Name are both required.")
            continue

        if code in seen_codes:
            result.errors.append(f"Row {row_number}: ZT Code '{code}' appears twice in this file.")
            continue
        seen_codes.add(code)

        category_code = str(_clean(cell(row, "Category code"))).upper()
        location_code = str(_clean(cell(row, "Location code"))).upper()
        category = categories.get(category_code)
        location = locations.get(location_code)
        if not category:
            result.errors.append(f"Row {row_number}: no category with code '{category_code}'.")
            continue
        if location_code and not location:
            result.errors.append(f"Row {row_number}: no location with code '{location_code}'.")
            continue

        status_text = str(_clean(cell(row, "Status"))).lower()
        status = (
            STATUS_LOOKUP.get(status_text, EquipmentStatus.AVAILABLE)
            if status_text
            else EquipmentStatus.AVAILABLE
        )
        if status_text and status_text not in STATUS_LOOKUP:
            result.errors.append(
                f"Row {row_number}: status '{status_text}' is not valid. "
                "See the Reference sheet for accepted values."
            )
            continue
        if status == EquipmentStatus.IN_PRODUCTION:
            # In production follows the Products page (fitted to a running
            # product), never a spreadsheet cell.
            current = Equipment.objects.filter(code=code).first()
            if not (current and current.running_fitting()):
                status = EquipmentStatus.AVAILABLE

        values = {
            "name": name,
            "category": category,
            "location": location,
            "status": status,
            "manufacturer": str(_clean(cell(row, "Manufacturer"))),
            "model_number": str(_clean(cell(row, "Model"))),
            "serial_number": str(_clean(cell(row, "Serial number"))),
            "year_made": _parse_int(cell(row, "Year"), "Year", row_number, result.errors),
            "is_measuring_device": _is_truthy(cell(row, "Needs calibration (x)")),
            "specifications": str(_clean(cell(row, "Specifications"))),
            "notes": str(_clean(cell(row, "Notes"))),
            "asset_no": str(_clean(cell(row, "Asset No."))).upper(),
            "equipment_no": str(_clean(cell(row, "Equipment No"))).upper(),
        }

        equipment_no = values["equipment_no"]
        if equipment_no:
            if equipment_no in seen_equipment_nos:
                result.errors.append(
                    f"Row {row_number}: Equipment No '{equipment_no}' appears twice in this file."
                )
                continue
            seen_equipment_nos.add(equipment_no)
            holder = Equipment.objects.filter(equipment_no=equipment_no).exclude(code=code).first()
            if holder:
                result.errors.append(
                    f"Row {row_number}: Equipment No '{equipment_no}' is already on {holder.code}."
                )
                continue

        asset_no = values["asset_no"]
        if asset_no:
            if asset_no in seen_asset_nos:
                result.errors.append(f"Row {row_number}: Asset No. '{asset_no}' appears twice in this file.")
                continue
            seen_asset_nos.add(asset_no)
            owner = Equipment.objects.filter(asset_no=asset_no).exclude(code=code).first()
            if owner:
                result.errors.append(f"Row {row_number}: Asset No. '{asset_no}' is already on {owner.code}.")
                continue

        existing = Equipment.objects.filter(code=code).first()
        if existing:
            if not update_existing:
                result.skipped += 1
                continue
            for attribute, value in values.items():
                setattr(existing, attribute, value)
            pending_update.append(existing)
        else:
            pending_new.append(Equipment(code=code, **values))

    if result.errors:
        return result

    with transaction.atomic():
        for item in pending_new:
            item.save()
            result.created += 1
        for item in pending_update:
            item.save()
            result.updated += 1

    return result


def _embed_photo(sheet, cell_ref: str, attachment) -> None:
    """Drop ``attachment``'s thumbnail into the cell, shrunk to fit it."""
    if attachment is None or not attachment.thumbnail:
        return
    try:
        with attachment.thumbnail.open("rb") as fh:
            raw = fh.read()
    except (FileNotFoundError, OSError):
        return  # the file record exists but the file is missing on disk
    try:
        with PILImage.open(BytesIO(raw)) as source:
            small = source.convert("RGB")
            small.thumbnail((PHOTO_CELL_PX, PHOTO_CELL_PX))
            width, height = small.size
            buffer = BytesIO()
            small.save(buffer, format="JPEG", quality=80)
    except OSError:
        return  # not a readable image
    buffer.seek(0)
    image = XLImage(buffer)
    image.width, image.height = width, height
    sheet.add_image(image, cell_ref)


def export_equipment_to_xlsx(queryset) -> Workbook:
    """Export the current (filtered) equipment list, with each asset's main photo."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Equipment"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F6FEB")
    headers = [
        "Photo", "ZT Code", "Equipment No", "Asset No.", "Name", "Category", "Location", "Status",
        "Manufacturer", "Model", "Serial number", "Year",
        "Used for (products)", "Needs calibration", "Specifications", "Notes",
    ]
    for index, label in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = max(len(label) + 4, 16)
    sheet.column_dimensions["A"].width = PHOTO_CELL_PX / 7  # ~pixels-to-Excel-units
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "B2"

    row_index = 2
    for item in queryset.select_related(
        "category", "location", "supplier"
    ).prefetch_related("attachments"):
        sheet.append([
            None,  # the photo is dropped in as a floating image, not a cell value
            item.code,
            item.equipment_no,
            item.asset_no,
            item.name,
            item.category.name if item.category else "",
            item.location.name if item.location else "",
            item.get_status_display(),
            item.manufacturer,
            item.model_number,
            item.serial_number,
            item.year_made,
            ", ".join(item.product_codes),
            "x" if item.is_measuring_device else "",
            item.specifications,
            item.notes,
        ])
        _embed_photo(sheet, f"A{row_index}", item.main_photo)
        sheet.row_dimensions[row_index].height = PHOTO_CELL_PX * 0.75  # px -> points
        row_index += 1

    return workbook
