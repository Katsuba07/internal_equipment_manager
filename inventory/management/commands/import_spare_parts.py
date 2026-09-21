"""
Import the spare parts listed in the team's workbook.

Reads the "Spare Part Stock" sheet of *Overall Equipment Management.xlsx* —
by default the copy on the engineering share — and creates one stocked part per
row, with the picture that sits in the sheet's Image column.

Dry run unless ``--commit``:

    manage.py import_spare_parts
    manage.py import_spare_parts --commit --photos
    manage.py import_spare_parts --file "C:/somewhere/else.xlsx" --commit

Run it again later and it updates the same rows by part number; the opening
stock is only recorded the first time, so balances are never double-counted.
"""

import pathlib
import re
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

import openpyxl
from django.core.files.base import ContentFile
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from core.models import Supplier
from core.photos import save_photo
from inventory.models import InventoryItem, ItemKind, StockTransaction, TransactionType

DEFAULT_FILE = r"\\192.168.10.253\Engineer\TE\02. Equipment Management\Overall Equipment Management.xlsx"
DEFAULT_SHEET = "Spare Part Stock"

# Their "Category" column, in our vocabulary.
KIND = {
    "test probe": ItemKind.PROBE,
    "probe": ItemKind.PROBE,
    "cable": ItemKind.CABLE,
    "harness": ItemKind.CABLE,
    "tool": ItemKind.TOOL,
    "consumable": ItemKind.CONSUMABLE,
    "spare unit": ItemKind.SPARE_EQUIPMENT,
}

# Makers that carry no information.
UNKNOWN_MAKERS = {"", "unknow", "unknown", "n/a", "na", "-"}

EXCEL_EPOCH = datetime(1899, 12, 30)


def excel_date(value) -> date | None:
    """The sheet keeps dates as plain serial numbers."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return (EXCEL_EPOCH + timedelta(days=float(value))).date()
    except (TypeError, ValueError):
        return None


def decimal(value, default="0") -> Decimal:
    try:
        return Decimal(str(value if value not in (None, "") else default))
    except (InvalidOperation, ValueError):
        return Decimal(default)


def sheet_pictures(path: pathlib.Path, sheet_name: str) -> dict[int, bytes]:
    """
    The picture sitting on each row of the Image column.

    openpyxl drops floating images, so the file is read as the zip it is: the
    sheet's drawing gives each picture's anchor row, and the drawing's
    relationships give the file inside the workbook.
    """
    pictures: dict[int, bytes] = {}
    with zipfile.ZipFile(path) as archive:
        workbook = archive.read("xl/workbook.xml").decode("utf-8", "ignore")
        order = re.findall(r'<sheet name="([^"]+)"', workbook)
        if sheet_name not in order:
            return pictures
        sheet_file = f"xl/worksheets/sheet{order.index(sheet_name) + 1}.xml"
        try:
            rels = archive.read(sheet_file.replace("worksheets/", "worksheets/_rels/") + ".rels")
        except KeyError:
            return pictures
        targets = re.findall(r'Target="([^"]+)"', rels.decode("utf-8", "ignore"))
        drawings = [t for t in targets if "drawing" in t]
        if not drawings:
            return pictures
        drawing_file = "xl/" + drawings[0].replace("../", "")
        drawing = archive.read(drawing_file).decode("utf-8", "ignore")
        anchors = re.findall(
            r"<xdr:from>.*?<xdr:row>(\d+)</xdr:row>.*?r:embed=\"(rId\d+)\"", drawing, re.S
        )
        mapping = dict(re.findall(
            r'Id="(rId\d+)"[^>]*Target="([^"]+)"',
            archive.read(drawing_file.replace("drawings/", "drawings/_rels/") + ".rels").decode("utf-8", "ignore"),
        ))
        for row, embed in anchors:
            target = mapping.get(embed)
            if not target:
                continue
            media = "xl/" + target.replace("../", "")
            try:
                # The anchor row is zero-based; sheet rows are not.
                pictures[int(row) + 1] = archive.read(media)
            except KeyError:
                continue
    return pictures


class Command(BaseCommand):
    help = "Import the spare parts sheet from the engineering workbook (dry run unless --commit)."

    def add_arguments(self, parser):
        parser.add_argument("--file", default=DEFAULT_FILE, help="Path to the workbook.")
        parser.add_argument("--sheet", default=DEFAULT_SHEET, help="Sheet to read.")
        parser.add_argument("--commit", action="store_true", help="Write to the database.")
        parser.add_argument("--photos", action="store_true", help="Also attach the sheet's pictures.")
        parser.add_argument("--plant", default="", help="Plant these parts are stored in (CTV1/CTV2/CTVN).")

    def handle(self, *args, **options):
        path = pathlib.Path(options["file"])
        if not path.exists():
            raise CommandError(f"{path} not found — is the share reachable?")

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if options["sheet"] not in workbook.sheetnames:
            raise CommandError(f"No sheet named {options['sheet']!r}. Found: {workbook.sheetnames}")
        sheet = workbook[options["sheet"]]

        rows = list(sheet.iter_rows(values_only=True))
        header = [str(c or "").strip().lower() for c in rows[0]]
        index = {name: position for position, name in enumerate(header)}
        needed = ["spare part id", "spare part name", "unit", "stock qty"]
        missing = [name for name in needed if name not in index]
        if missing:
            raise CommandError(f"The sheet is missing these columns: {missing}")

        pictures = sheet_pictures(path, options["sheet"]) if options["photos"] else {}
        self.stdout.write(f"Rows: {len(rows) - 1} | pictures found: {len(pictures)}")

        created = updated = photos = 0
        with transaction.atomic():
            for number, row in enumerate(rows[1:], start=2):
                def cell(name):
                    position = index.get(name)
                    return row[position] if position is not None and position < len(row) else None

                code = str(cell("spare part id") or "").strip().upper()
                if not code:
                    continue

                maker = str(cell("manufacturer") or "").strip()
                supplier = None
                if maker.lower() not in UNKNOWN_MAKERS:
                    supplier = Supplier.objects.filter(name__iexact=maker).first()
                    if supplier is None:
                        supplier = Supplier.objects.create(
                            code=re.sub(r"[^A-Z0-9]+", "-", maker.upper())[:30] or "SUP",
                            name=maker[:150],
                        )

                total = decimal(cell("total qty"))
                on_shelf = decimal(cell("stock qty"))
                in_use = total - on_shelf
                notes = [f"From {path.name}, sheet “{options['sheet']}”, row {number}."]
                if in_use > 0:
                    notes.append(f"Workbook total {total:g}, of which {in_use:g} already in use.")
                if maker and maker.lower() not in UNKNOWN_MAKERS:
                    notes.append(f"Manufacturer: {maker}")
                for label, column in (("Last restock", "last restock date"), ("Next restock", "next restock date")):
                    when = excel_date(cell(column))
                    if when:
                        notes.append(f"{label}: {when:%d %b %Y}")

                fields = {
                    "name": str(cell("spare part name") or code)[:200],
                    "kind": KIND.get(str(cell("category") or "").strip().lower(), ItemKind.SPARE_PART),
                    "unit": str(cell("unit") or "pcs")[:20],
                    "specification": str(cell("category") or "")[:255],
                    "min_quantity": decimal(cell("min stock level")),
                    "supplier": supplier,
                    "plant": options["plant"],
                    "notes": "\n".join(notes),
                }

                item = InventoryItem.objects.filter(code=code).first()
                if item:
                    for key, value in fields.items():
                        setattr(item, key, value)
                    item.save()
                    updated += 1
                else:
                    item = InventoryItem.objects.create(code=code, **fields)
                    created += 1
                    if on_shelf:
                        StockTransaction(
                            item=item,
                            transaction_type=TransactionType.IN,
                            quantity=on_shelf,
                            reason="Opening stock from the engineering workbook",
                        ).apply()

                picture = pictures.get(number)
                if picture and options["commit"] and not item.photos:
                    try:
                        save_photo(item, ContentFile(picture, name=f"{code}.png"), None)
                        photos += 1
                    except Exception as error:      # one odd image must not stop the run
                        self.stderr.write(f"{code}: {error}")

            if not options["commit"]:
                transaction.set_rollback(True)

        mode = "WROTE" if options["commit"] else "DRY RUN — nothing was saved"
        self.stdout.write(self.style.SUCCESS(
            f"{mode}: {created} created, {updated} updated, {photos} photos"
        ))
